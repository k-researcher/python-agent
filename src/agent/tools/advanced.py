from __future__ import annotations

import asyncio
import json
import zipfile
from pathlib import Path
from typing import Any
from urllib.parse import quote_plus

import asyncssh
import httpx
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import create_async_engine

from agent.config import Settings
from agent.security import PathGuard
from agent.tools.base import RiskLevel, Tool, ToolContext, ToolError
from agent.tools.network import validate_outbound_url

MAX_WORKBOOK_BYTES = 50_000_000
MAX_WORKBOOK_UNCOMPRESSED_BYTES = 200_000_000
MAX_WORKBOOK_PARTS = 5000
FORBIDDEN_HTTP_HEADERS = {
    "connection",
    "content-length",
    "host",
    "proxy-authorization",
    "te",
    "transfer-encoding",
    "upgrade",
}


async def audit(
    context: ToolContext,
    category: str,
    destination: str,
    operation: str,
    payload: str,
    status: str,
    detail: str | None = None,
) -> None:
    if context.audit_egress is not None:
        await context.audit_egress(category, destination, operation, payload, status, detail)


class HttpTool(Tool):
    name = "http_request"
    description = "Send an approved HTTP request to a host from the configured allowlist."
    risk_level = RiskLevel.network_access
    network_capability = True
    input_schema = {
        "type": "object",
        "properties": {
            "url": {"type": "string"},
            "method": {"type": "string", "enum": ["GET", "POST", "PUT", "PATCH", "DELETE"]},
            "headers": {"type": "object", "additionalProperties": {"type": "string"}},
            "body": {"type": ["object", "array", "string", "null"]},
        },
        "required": ["url"],
        "additionalProperties": False,
    }

    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    async def execute(self, context: ToolContext, arguments: dict[str, Any]) -> dict[str, Any]:
        url = str(arguments["url"])
        await validate_outbound_url(url, self.settings.network_allowlist)
        method = str(arguments.get("method", "GET")).upper()
        headers = {str(k): str(v) for k, v in dict(arguments.get("headers") or {}).items()}
        forbidden = sorted(name for name in headers if name.lower() in FORBIDDEN_HTTP_HEADERS)
        if forbidden:
            raise ToolError(f"Forbidden HTTP headers: {', '.join(forbidden)}")
        body = arguments.get("body")
        payload = "" if body is None else json.dumps(body, ensure_ascii=False)
        await audit(context, "http", url, method, payload, "started")

        request_kwargs: dict[str, Any] = {"headers": headers}
        if body is not None:
            request_kwargs["content"] = payload.encode()
            headers.setdefault("Content-Type", "application/json")

        collected = bytearray()
        try:
            async with (
                httpx.AsyncClient(timeout=30, follow_redirects=False) as client,
                client.stream(method, url, **request_kwargs) as response,
            ):
                async for chunk in response.aiter_bytes():
                    collected.extend(chunk)
                    if len(collected) > self.settings.http_max_response_bytes:
                        raise ToolError("HTTP response exceeded the configured size limit")
                result = {
                    "status": response.status_code,
                    "content_type": response.headers.get("content-type", ""),
                    "location": response.headers.get("location"),
                    "body": collected.decode(errors="replace"),
                }
            await audit(context, "http", url, method, payload, "completed")
            return result
        except Exception as exc:
            await audit(context, "http", url, method, payload, "error", str(exc))
            raise


class WebSearchTool(Tool):
    name = "web_search"
    description = "Search the web through the configured search endpoint."
    risk_level = RiskLevel.network_access
    network_capability = True
    input_schema = {
        "type": "object",
        "properties": {"query": {"type": "string"}},
        "required": ["query"],
        "additionalProperties": False,
    }

    def __init__(self, settings: Settings) -> None:
        self.settings = settings
        self.http = HttpTool(settings)

    async def execute(self, context: ToolContext, arguments: dict[str, Any]) -> dict[str, Any]:
        template = self.settings.web_search_url_template
        if not template or "{query}" not in template:
            raise ToolError("AGENT_WEB_SEARCH_URL_TEMPLATE is not configured")
        query = str(arguments["query"])
        url = template.replace("{query}", quote_plus(query))
        return await self.http.execute(context, {"url": url, "method": "GET"})


def serialize_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    primitives = (str, int, float, bool, type(None))
    return [
        {
            key: value
            if isinstance(value, primitives) and not isinstance(value, str)
            else str(value)[:10_000]
            for key, value in row.items()
        }
        for row in rows
    ]


class DbSelectTool(Tool):
    name = "db_select"
    description = "Execute one approved read-only SQL query on a configured connection."
    risk_level = RiskLevel.network_access
    network_capability = True
    input_schema = {
        "type": "object",
        "properties": {
            "connection": {"type": "string"},
            "query": {"type": "string"},
            "parameters": {"type": "object"},
            "limit": {"type": "integer", "minimum": 1, "maximum": 500, "default": 100},
        },
        "required": ["connection", "query"],
        "additionalProperties": False,
    }

    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    def connection_url(self, name: str) -> str:
        url = self.settings.database_connections.get(name)
        if not url:
            raise ToolError(f"Unknown database connection: {name}")
        return url

    async def execute(self, context: ToolContext, arguments: dict[str, Any]) -> dict[str, Any]:
        name = str(arguments["connection"])
        query = str(arguments["query"]).strip()
        normalized = query.rstrip(";").lstrip().lower()
        if ";" in normalized or not normalized.startswith(("select ", "with ", "explain ")):
            raise ToolError("db_select accepts one SELECT/WITH/EXPLAIN statement")
        url = self.connection_url(name)
        await audit(context, "database", name, "select", query, "started")
        dialect = make_url(url).get_backend_name()
        if dialect not in {"sqlite", "postgresql"}:
            raise ToolError("db_select supports enforced read-only mode for SQLite/PostgreSQL")
        engine = create_async_engine(url, pool_pre_ping=True)
        try:
            async with engine.begin() as connection:
                if dialect == "postgresql":
                    await connection.execute(text("SET TRANSACTION READ ONLY"))
                else:
                    await connection.execute(text("PRAGMA query_only = ON"))
                result = await connection.execute(
                    text(query), dict(arguments.get("parameters") or {})
                )
                limit = min(500, int(arguments.get("limit", 100)))
                rows = [dict(row) for row in result.mappings().fetchmany(limit)]
            await audit(context, "database", name, "select", query, "completed")
            return {"rows": serialize_rows(rows), "count": len(rows), "limit": limit}
        except Exception as exc:
            await audit(context, "database", name, "select", query, "error", str(exc))
            raise
        finally:
            await engine.dispose()


class DbExecuteTool(DbSelectTool):
    name = "db_execute"
    description = "Execute one approved mutating SQL statement on a configured connection."
    risk_level = RiskLevel.remote_write

    async def execute(self, context: ToolContext, arguments: dict[str, Any]) -> dict[str, Any]:
        name = str(arguments["connection"])
        query = str(arguments["query"]).strip()
        if ";" in query.rstrip(";"):
            raise ToolError("Only one SQL statement is allowed")
        if query.lstrip().lower().startswith(("select ", "with ", "explain ")):
            raise ToolError("Use db_select for read-only queries")
        url = self.connection_url(name)
        await audit(context, "database", name, "execute", query, "started")
        engine = create_async_engine(url, pool_pre_ping=True)
        try:
            async with engine.begin() as connection:
                result = await connection.execute(
                    text(query), dict(arguments.get("parameters") or {})
                )
            await audit(context, "database", name, "execute", query, "completed")
            return {"rowcount": result.rowcount}
        except Exception as exc:
            await audit(context, "database", name, "execute", query, "error", str(exc))
            raise
        finally:
            await engine.dispose()


class SshExecTool(Tool):
    name = "ssh_exec"
    description = "Execute an approved command on a configured SSH connection."
    risk_level = RiskLevel.remote_write
    network_capability = True
    input_schema = {
        "type": "object",
        "properties": {
            "connection": {"type": "string"},
            "command": {"type": "string"},
            "timeout_seconds": {"type": "integer", "minimum": 1, "maximum": 300, "default": 60},
        },
        "required": ["connection", "command"],
        "additionalProperties": False,
    }

    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    async def execute(self, context: ToolContext, arguments: dict[str, Any]) -> dict[str, Any]:
        name = str(arguments["connection"])
        config = self.settings.ssh_connections.get(name)
        if not config:
            raise ToolError(f"Unknown SSH connection: {name}")
        if not config.get("known_hosts"):
            raise ToolError("SSH connection must configure known_hosts")
        command = str(arguments["command"])
        await audit(context, "ssh", name, "exec", command, "started")
        options: dict[str, Any] = {
            "host": config["host"],
            "port": int(config.get("port", 22)),
            "username": config.get("username"),
            "known_hosts": config["known_hosts"],
        }
        if config.get("password"):
            options["password"] = config["password"]
        if config.get("client_keys"):
            options["client_keys"] = config["client_keys"]
        try:
            async with asyncssh.connect(**options) as connection:
                result = await connection.run(
                    command, check=False, timeout=int(arguments.get("timeout_seconds", 60))
                )
            await audit(context, "ssh", name, "exec", command, "completed")
            return {
                "exit_code": result.exit_status,
                "stdout": str(result.stdout)[:100_000],
                "stderr": str(result.stderr)[:100_000],
            }
        except Exception as exc:
            await audit(context, "ssh", name, "exec", command, "error", str(exc))
            raise


class ExcelSheetsTool(Tool):
    name = "excel_sheets"
    description = "List worksheet names in an XLSX workbook inside the project."
    risk_level = RiskLevel.read_only
    input_schema = {
        "type": "object",
        "properties": {"path": {"type": "string"}},
        "required": ["path"],
        "additionalProperties": False,
    }

    async def execute(self, context: ToolContext, arguments: dict[str, Any]) -> dict[str, Any]:
        path = PathGuard(context.project_root).resolve(str(arguments["path"]))
        self.validate_workbook(path)
        return await asyncio.to_thread(self.read_sheets, path)

    @staticmethod
    def validate_workbook(path: Path) -> None:
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ToolError("Only XLSX/XLSM files are supported")
        if path.stat().st_size > MAX_WORKBOOK_BYTES:
            raise ToolError("Workbook exceeds the compressed size limit")
        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > MAX_WORKBOOK_PARTS:
                    raise ToolError("Workbook contains too many archive parts")
                if sum(entry.file_size for entry in entries) > MAX_WORKBOOK_UNCOMPRESSED_BYTES:
                    raise ToolError("Workbook exceeds the uncompressed size limit")
        except zipfile.BadZipFile as exc:
            raise ToolError("Workbook is not a valid XLSX archive") from exc

    @staticmethod
    def read_sheets(path: Path) -> dict[str, Any]:
        workbook = load_workbook(
            path, read_only=True, data_only=True, keep_links=False
        )
        try:
            return {"sheets": workbook.sheetnames}
        finally:
            workbook.close()


class ExcelReadTool(ExcelSheetsTool):
    name = "excel_read"
    description = "Read a bounded cell range from an XLSX worksheet."
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "sheet": {"type": "string"},
            "min_row": {"type": "integer", "minimum": 1, "default": 1},
            "max_row": {"type": "integer", "minimum": 1, "default": 100},
            "min_col": {"type": "integer", "minimum": 1, "default": 1},
            "max_col": {"type": "integer", "minimum": 1, "default": 30},
        },
        "required": ["path", "sheet"],
        "additionalProperties": False,
    }

    async def execute(self, context: ToolContext, arguments: dict[str, Any]) -> dict[str, Any]:
        path = PathGuard(context.project_root).resolve(str(arguments["path"]))
        self.validate_workbook(path)
        return await asyncio.to_thread(self.read_range, path, arguments)

    @staticmethod
    def read_range(path: Path, arguments: dict[str, Any]) -> dict[str, Any]:
        workbook = load_workbook(
            path, read_only=True, data_only=True, keep_links=False
        )
        try:
            sheet_name = str(arguments["sheet"])
            if sheet_name not in workbook.sheetnames:
                raise ToolError(f"Worksheet not found: {sheet_name}")
            min_row = int(arguments.get("min_row", 1))
            max_row = min(min_row + 499, int(arguments.get("max_row", 100)))
            min_col = int(arguments.get("min_col", 1))
            max_col = min(min_col + 99, int(arguments.get("max_col", 30)))
            rows = [
                [cell.value for cell in row]
                for row in workbook[sheet_name].iter_rows(
                    min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
                )
            ]
            return {
                "sheet": sheet_name,
                "range": [min_row, max_row, min_col, max_col],
                "rows": rows,
            }
        finally:
            workbook.close()


class RunAgentTool(Tool):
    name = "run_agent"
    description = "Create a bounded child agent session and optionally wait for its final response."
    risk_level = RiskLevel.network_access
    allowed_modes = frozenset({"dev"})
    def __init__(self, settings: Settings) -> None:
        self.input_schema = {
            "type": "object",
            "properties": {
                "mode": {"type": "string", "enum": ["dev", "ask"]},
                "prompt": {"type": "string"},
                "wait": {"type": "boolean", "default": True},
                "llm_profile": {
                    "type": "string",
                    "enum": sorted(settings.available_llm_profiles()),
                    "description": "Optional child model profile; inherits parent when omitted.",
                },
            },
            "required": ["mode", "prompt"],
            "additionalProperties": False,
        }

    async def execute(self, context: ToolContext, arguments: dict[str, Any]) -> dict[str, Any]:
        if context.spawn_child is None:
            raise ToolError("Child agent runtime is unavailable")
        return await context.spawn_child(
            str(arguments["mode"]),
            str(arguments["prompt"]),
            bool(arguments.get("wait", True)),
            str(arguments["llm_profile"]) if arguments.get("llm_profile") else None,
        )


def advanced_tools(settings: Settings) -> tuple[Tool, ...]:
    return (
        HttpTool(settings),
        WebSearchTool(settings),
        DbSelectTool(settings),
        DbExecuteTool(settings),
        SshExecTool(settings),
        ExcelSheetsTool(),
        ExcelReadTool(),
        RunAgentTool(settings),
    )
